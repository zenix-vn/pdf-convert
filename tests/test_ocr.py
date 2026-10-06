import docx
import numpy as np
import pymupdf
import pytest

from zenix_pdf.core.converter import PdfConverterService, _normalize_ocr_text
from zenix_pdf.core.models import ConvertJob, ConvertOptions
from zenix_pdf.core.ocr import (
    _fit_islands,
    _row_runs,
    _split_words,
    find_tessdata,
    page_needs_ocr,
)
from zenix_pdf.core.pdf_info import inspect_pdf

needs_tessdata = pytest.mark.skipif(
    find_tessdata() is None, reason="chưa có dữ liệu OCR (python tools/fetch_tessdata.py)"
)

LINES = [
    "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
    "Họ và tên người khám sức khỏe",
    "Nơi ở hiện tại: Cầu Giấy, Hà Nội",
]


def _docx_text(path) -> str:
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]

    def walk(tables):
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    parts.extend(p.text for p in cell.paragraphs)
                    walk(cell.tables)

    walk(d.tables)
    return "\n".join(parts)


def _text_page_pdf(path):
    """Trang chữ thật (Times New Roman) để render thành ảnh."""
    doc = pymupdf.open()
    page = doc.new_page()
    font = pymupdf.Font(fontfile=r"C:\Windows\Fonts\times.ttf")
    tw = pymupdf.TextWriter(page.rect)
    for i, line in enumerate(LINES):
        tw.append((72, 100 + i * 30), line, font=font, fontsize=14)
    tw.write_text(page)
    page.draw_rect(pymupdf.Rect(72, 220, 400, 300), width=1)  # khung bảng
    page.draw_line((236, 220), (236, 300), width=1)
    doc.save(str(path))
    return path


@pytest.fixture
def scanned_pdf(tmp_path):
    """PDF chỉ chứa ảnh chụp của trang chữ."""
    src = pymupdf.open(str(_text_page_pdf(tmp_path / "src.pdf")))
    pix = src[0].get_pixmap(dpi=200)
    out = pymupdf.open()
    page = out.new_page(width=src[0].rect.width, height=src[0].rect.height)
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    path = tmp_path / "scan.pdf"
    out.save(str(path))
    return path


def test_page_needs_ocr(make_pdf, scanned_pdf):
    with pymupdf.open(str(make_pdf(pages=1))) as doc:
        assert not page_needs_ocr(doc[0])
    with pymupdf.open(str(scanned_pdf)) as doc:
        assert page_needs_ocr(doc[0])
    assert inspect_pdf(scanned_pdf).is_scanned


def test_page_needs_ocr_outlined_text(tmp_path):
    """Chữ bị convert thành đường vector (CorelDRAW/Illustrator): không có lớp text."""
    doc = pymupdf.open()
    page = doc.new_page()
    shape = page.new_shape()
    for i in range(1200):
        x, y = 50 + (i % 60) * 8, 80 + (i // 60) * 20
        shape.draw_bezier((x, y), (x + 2, y - 4), (x + 4, y - 4), (x + 6, y))
        shape.draw_line((x, y), (x + 6, y))
    shape.finish(fill=(0, 0, 0))
    shape.commit()
    assert page_needs_ocr(page)


def test_split_words_restores_missing_space():
    # MuPDF làm mất dấu cách sau chữ có dấu: "Họ" và "tên" sát nhau nhưng có khe
    chars = [{"c": c, "bbox": (x, 0, x + 5, 10)} for c, x in
             [("H", 0), ("ọ", 5), ("t", 14), ("ê", 19), ("n", 24)]]
    assert [w.text for w in _split_words(chars)] == ["Họ", "tên"]


def test_split_words_combining_marks():
    chars = [{"c": c, "bbox": (x, 0, x + 5, 10)} for c, x in
             [("T", 0), ("Y", 5), ("\u0300", 10)]]
    assert [w.text for w in _split_words(chars)] == ["TỲ"]


def test_fit_islands():
    cum = np.cumsum([0.0, 10, 10, 10, 10, 10])  # 5 ký tự rộng 10
    cost, cuts = _fit_islands(cum, [(0, 20), (30, 60)])
    assert cuts == [(0, 2), (2, 5)] and cost == 0


def test_row_runs_keeps_diacritics_with_line():
    own = np.zeros((60, 50), bool)
    own[0:3, 10:14] = True  # dấu thanh phía trên dòng 1
    own[6:20, :] = True  # dòng 1
    own[34:48, :] = True  # dòng 2
    own[50:53, 20:23] = True  # dấu nặng dưới dòng 2
    runs = _row_runs(own)
    assert len(runs) == 2
    assert runs[0][0] == 0 and runs[1][1] == 60


def test_normalize_ocr_text(tmp_path):
    path = tmp_path / "a.docx"
    d = docx.Document()
    d.add_paragraph("GKSK\u00adTYT\u00a0x")
    d.save(str(path))
    _normalize_ocr_text(path)
    assert docx.Document(str(path)).paragraphs[0].text == "GKSK-TYT x"


@needs_tessdata
def test_convert_scanned_pdf(scanned_pdf):
    job = ConvertJob(pdf_path=scanned_pdf, page_count=1)
    out = PdfConverterService().convert(job, ConvertOptions())
    text = _docx_text(out)
    for word in ["CỘNG", "VIỆT", "Họ và tên", "khám sức khỏe", "Cầu Giấy"]:
        assert word in text, (word, text)
    assert docx.Document(str(out)).tables  # khung vẽ trong ảnh -> bảng thật


def test_convert_scanned_without_ocr(scanned_pdf):
    job = ConvertJob(pdf_path=scanned_pdf, page_count=1)
    out = PdfConverterService().convert(job, ConvertOptions(ocr=False))
    assert "Cầu Giấy" not in _docx_text(out)
